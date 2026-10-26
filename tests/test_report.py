import io
import json

from docx import Document

from auditor.analyzer import LegalAnalyzer
from auditor.extractor import extract_pdf
from auditor.report import DISCLAIMER, to_docx, to_json, to_markdown
from tests.conftest import FakeClient, fake_response


def _outcome(settings, pdf, payload):
    return LegalAnalyzer(settings, client=FakeClient(fake_response(payload))).analyze(extract_pdf(pdf), pdf_bytes=pdf)


def test_markdown_report(settings, sample_pdf_bytes, sample_payload):
    md = to_markdown(_outcome(settings, sample_pdf_bytes, sample_payload), "doc.pdf")
    assert "## Nivel de riesgo: CRÍTICO (82/100)" in md
    assert md.index("Hipoteca vigente") < md.index("Diferencia de superficie")  # sorted by severity
    assert "No se identifica en el documento" in md
    assert DISCLAIMER in md


def test_docx_report_opens(settings, sample_pdf_bytes, sample_payload):
    data = to_docx(_outcome(settings, sample_pdf_bytes, sample_payload), "doc.pdf")
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "NIVEL DE RIESGO: CRÍTICO" in text
    assert "Hipoteca vigente" in text
    assert len(doc.tables) == 2


def test_json_report(settings, sample_pdf_bytes, sample_payload):
    payload = json.loads(to_json(_outcome(settings, sample_pdf_bytes, sample_payload), "doc.pdf"))
    assert payload["analisis"]["puntuacion_riesgo"] == 82
    assert payload["tokens"] == {"entrada": 1234, "salida": 567}
