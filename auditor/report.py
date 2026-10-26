"""Deliverable rendering: Markdown, Word (.docx) and JSON."""

from __future__ import annotations

import io
import json
from datetime import datetime
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from auditor.analyzer import AnalysisOutcome
from auditor.schemas import CATEGORY_LABELS, RISK_LABELS, SEVERITY_LABELS, Finding, RiskReport

TIMEZONE = ZoneInfo("America/Mexico_City")
DISCLAIMER = (
    "Este reporte fue generado con asistencia de inteligencia artificial a partir exclusivamente del documento "
    "proporcionado. No constituye opinión legal ni sustituye la revisión de un notario, abogado o perito, ni la "
    "consulta directa al Registro Público de la Propiedad y a las autoridades de desarrollo urbano."
)
_NA = "No se identifica en el documento"

RISK_COLORS = {"CRITICO": "B42318", "ALTO": "DC6803", "MEDIO": "B54708", "BAJO": "067647"}


def _now() -> str:
    return datetime.now(TIMEZONE).strftime("%d/%m/%Y %H:%M")


def _v(value: str | None) -> str:
    return value if value else _NA


def property_rows(report: RiskReport) -> list[tuple[str, str]]:
    p = report.inmueble
    return [
        ("Ubicación", _v(p.ubicacion)),
        ("Clave catastral", _v(p.clave_catastral)),
        ("Folio real", _v(p.folio_real)),
        ("Superficie de terreno", _v(p.superficie_terreno)),
        ("Superficie de construcción", _v(p.superficie_construccion)),
        ("Titular", _v(p.titular)),
    ]


def zoning_rows(report: RiskReport) -> list[tuple[str, str]]:
    z = report.zonificacion
    return [
        ("Uso de suelo", _v(z.uso_de_suelo)),
        ("Niveles máximos", _v(z.niveles_maximos)),
        ("Altura máxima", _v(z.altura_maxima)),
        ("Área libre mínima", _v(z.area_libre_minima)),
        ("COS", _v(z.cos)),
        ("CUS", _v(z.cus)),
        ("Densidad", _v(z.densidad)),
        ("Usos permitidos", ", ".join(z.usos_permitidos) or _NA),
        ("Usos prohibidos", ", ".join(z.usos_prohibidos) or _NA),
        ("Instrumento aplicable", _v(z.instrumento_aplicable)),
    ]


def _cell(value: str) -> str:
    """Make a value safe for a Markdown table cell."""
    return value.replace("|", "\\|").replace("\n", " ")


def _page(f: Finding) -> str:
    return f"p. {f.pagina}" if f.pagina else "página no indicada"


# --------------------------------------------------------------------- Markdown
def to_markdown(outcome: AnalysisOutcome, file_name: str) -> str:
    r = outcome.report
    md: list[str] = [
        "# Reporte de Due Diligence Legal-Urbana",
        f"**Documento:** {file_name}  \n**Tipo:** {r.tipo_documento}  \n**Fecha:** {_now()}",
        "",
        f"## Nivel de riesgo: {RISK_LABELS[r.nivel_riesgo_global]} ({r.puntuacion_riesgo}/100)",
        f"> {r.veredicto}",
        "",
        "## Resumen ejecutivo",
        r.resumen_ejecutivo,
        "",
        "## Identificación del inmueble",
        "| Campo | Valor |",
        "|---|---|",
        *[f"| {k} | {_cell(v)} |" for k, v in property_rows(r)],
        "",
        "## Zonificación",
        "| Parámetro | Valor |",
        "|---|---|",
        *[f"| {k} | {_cell(v)} |" for k, v in zoning_rows(r)],
        "",
        "## Hallazgos",
    ]
    findings = r.sorted_findings()
    if not findings:
        md.append("No se identificaron hallazgos de riesgo en el documento.")
    for i, f in enumerate(findings, start=1):
        md += [
            f"### {i}. [{SEVERITY_LABELS[f.severidad]}] {f.titulo}",
            f"*{CATEGORY_LABELS[f.categoria]} · {_page(f)}*",
            "",
            f.descripcion,
            "",
            f"> **Evidencia:** {f.evidencia}",
            "",
            f"**Recomendación:** {f.recomendacion}",
            "",
        ]
    for title, items in (
        ("Documentos a solicitar", r.documentos_faltantes),
        ("Siguientes pasos", r.siguientes_pasos),
        ("Limitaciones del análisis", r.limitaciones + outcome.warnings),
    ):
        if items:
            md.append(f"## {title}")
            md += [f"- {item}" for item in items]
            md.append("")
    md += ["---", f"*{DISCLAIMER}*", f"*Modelo: {outcome.model} · Modo: {outcome.mode}*"]
    return "\n".join(md)


# ------------------------------------------------------------------------- Word
def _table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for run in cells[0].paragraphs[0].runs:
            run.bold = True


def to_docx(outcome: AnalysisOutcome, file_name: str) -> bytes:
    r = outcome.report
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    doc.add_heading("Reporte de Due Diligence Legal-Urbana", level=0)
    meta = doc.add_paragraph()
    meta.add_run("Documento: ").bold = True
    meta.add_run(f"{file_name}\n")
    meta.add_run("Tipo: ").bold = True
    meta.add_run(f"{r.tipo_documento}\n")
    meta.add_run("Fecha: ").bold = True
    meta.add_run(_now())

    risk = doc.add_paragraph()
    risk.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = risk.add_run(f"NIVEL DE RIESGO: {RISK_LABELS[r.nivel_riesgo_global]}  ·  {r.puntuacion_riesgo}/100")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(RISK_COLORS[r.nivel_riesgo_global])
    verdict = doc.add_paragraph(r.veredicto)
    verdict.alignment = WD_ALIGN_PARAGRAPH.CENTER
    verdict.runs[0].italic = True

    doc.add_heading("Resumen ejecutivo", level=1)
    for paragraph in [p for p in r.resumen_ejecutivo.split("\n") if p.strip()]:
        doc.add_paragraph(paragraph.strip())

    doc.add_heading("Identificación del inmueble", level=1)
    _table(doc, property_rows(r))
    doc.add_heading("Zonificación", level=1)
    _table(doc, zoning_rows(r))

    doc.add_heading("Hallazgos", level=1)
    findings = r.sorted_findings()
    if not findings:
        doc.add_paragraph("No se identificaron hallazgos de riesgo en el documento.")
    for i, f in enumerate(findings, start=1):
        doc.add_heading(f"{i}. [{SEVERITY_LABELS[f.severidad]}] {f.titulo}", level=2)
        sub = doc.add_paragraph(f"{CATEGORY_LABELS[f.categoria]} · {_page(f)}")
        sub.runs[0].italic = True
        doc.add_paragraph(f.descripcion)
        ev = doc.add_paragraph()
        ev.add_run("Evidencia: ").bold = True
        ev.add_run(f.evidencia).italic = True
        rec = doc.add_paragraph()
        rec.add_run("Recomendación: ").bold = True
        rec.add_run(f.recomendacion)

    for title, items in (
        ("Documentos a solicitar", r.documentos_faltantes),
        ("Siguientes pasos", r.siguientes_pasos),
        ("Limitaciones del análisis", r.limitaciones + outcome.warnings),
    ):
        if items:
            doc.add_heading(title, level=1)
            for item in items:
                doc.add_paragraph(item, style="List Bullet")

    footer = doc.add_paragraph()
    footer_run = footer.add_run(f"{DISCLAIMER}\nModelo: {outcome.model} · Modo: {outcome.mode}")
    footer_run.italic = True
    footer_run.font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------------- JSON
def to_json(outcome: AnalysisOutcome, file_name: str) -> str:
    payload = {
        "documento": file_name,
        "generado": _now(),
        "modelo": outcome.model,
        "modo": outcome.mode,
        "paginas_analizadas": outcome.pages_analyzed,
        "tokens": {"entrada": outcome.input_tokens, "salida": outcome.output_tokens},
        "advertencias": outcome.warnings,
        "analisis": outcome.report.model_dump(),
    }
    return json.dumps(payload, ensure_ascii=False, indent=2)
